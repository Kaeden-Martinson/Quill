"""Word (.docx) import and export."""
import io
import zipfile

import docx
from docx.shared import Pt

MAX_UPLOAD = 5 * 1024 * 1024        # 5 MB on the wire
MAX_UNCOMPRESSED = 50 * 1024 * 1024  # guards against zip bombs


class BadFile(ValueError):
    pass


def read_docx(data):
    """Returns (title, body) with paragraphs separated by blank lines."""
    if len(data) > MAX_UPLOAD:
        raise BadFile("That file is over 5 MB.")
    not_docx = "That doesn't look like a Word (.docx) file. Older .doc files need saving as .docx first."
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            unpacked = sum(i.file_size for i in zf.infolist())
    except zipfile.BadZipFile:
        raise BadFile(not_docx)
    if unpacked > MAX_UNCOMPRESSED:
        raise BadFile("That file is too large once unpacked.")
    try:
        document = docx.Document(io.BytesIO(data))
    except (KeyError, ValueError):
        raise BadFile(not_docx)

    title = (document.core_properties.title or "").strip()
    paragraphs = []
    for p in document.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        # The first heading is the title (or repeats it), not part of the essay body.
        if not paragraphs and p.style.name.lower().startswith(("title", "heading")) and (not title or text == title):
            title = text
            continue
        paragraphs.append(text)
    return title, "\n\n".join(paragraphs)


def write_docx(title, body):
    document = docx.Document()
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    document.core_properties.title = title
    document.add_heading(title, level=1)
    for paragraph in body.split("\n\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph.strip().replace("\n", " "))
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
